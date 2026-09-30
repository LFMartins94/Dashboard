"""Gera fixtures anonimizadas a partir da planilha CAP real.

Os valores e a cardinalidade são preservados para manter o caso contábil de
aceitação. Identificadores, contas, histórico e nome da empresa são trocados.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook


RAIZ = Path(__file__).resolve().parents[1]
ORIGEM = Path.home() / "Downloads" / "CAP_ILHAS_DO_LAGO_CONCILIADO (1).xlsx"
DESTINO = RAIZ / "tests" / "fixtures" / "aceitacao"
DATA_FIXA = datetime(2026, 1, 1, 0, 0, 0)


def preparar_livro(livro: Workbook) -> None:
    livro.properties.creator = "ContaView"
    livro.properties.lastModifiedBy = "ContaView"
    livro.properties.created = DATA_FIXA
    livro.properties.modified = DATA_FIXA


def estabilizar_xlsx(caminho: Path) -> None:
    temporario = caminho.with_suffix(".deterministico.tmp")
    with ZipFile(caminho, "r") as origem, ZipFile(
        temporario,
        "w",
        compression=ZIP_DEFLATED,
        compresslevel=9,
    ) as destino:
        for nome in sorted(origem.namelist()):
            info = ZipInfo(nome, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            conteudo = origem.read(nome)
            if nome == "docProps/core.xml":
                conteudo = re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    rb'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" '
                    rb'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
                    rb'xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>',
                    conteudo,
                )
            destino.writestr(info, conteudo)
    temporario.replace(caminho)


def ler_origem() -> list[dict[str, object]]:
    if not ORIGEM.exists():
        raise SystemExit(f"Planilha CAP não encontrada: {ORIGEM}")
    planilha = load_workbook(ORIGEM, read_only=True, data_only=True)
    aba = planilha.active
    celulas = [str(linha[0]) for linha in aba.iter_rows(values_only=True) if linha and linha[0]]
    registros: list[dict[str, object]] = []
    contas: dict[str, str] = {}
    for numero, texto in enumerate(celulas[1:], start=1):
        partes = texto.split(";")
        if len(partes) < 6:
            raise SystemExit(f"Linha {numero + 1} da CAP possui menos de seis campos")
        conta_original = partes[1]
        if conta_original not in contas:
            contas[conta_original] = f"{len(contas) + 1}00000001"
        registros.append(
            {
                "data": datetime.strptime(partes[0], "%d/%m/%Y").date(),
                "conta": contas[conta_original],
                "valor": Decimal(partes[2].replace(".", "").replace(",", ".")),
                "tipo": partes[3],
                "historico": f"MOVIMENTO ANONIMIZADO {numero:03d}",
                "filial": "1",
            }
        )
    return registros


def valor_br(valor: Decimal) -> str:
    return f"{valor:.2f}".replace(".", ",")


def criar_xlsx_delimitado(registros: list[dict[str, object]], caminho: Path) -> None:
    livro = Workbook()
    preparar_livro(livro)
    aba = livro.active
    aba.title = "MOVIMENTOS_ANONIMIZADOS"
    aba.append(["Data;Conta Contábil;Valor;Tipo;Histórico;Filial"])
    for item in registros:
        aba.append(
            [
                ";".join(
                    [
                        item["data"].strftime("%d/%m/%Y"),
                        str(item["conta"]),
                        valor_br(item["valor"]),
                        str(item["tipo"]),
                        str(item["historico"]),
                        str(item["filial"]),
                    ]
                )
            ]
        )
    livro.save(caminho)
    estabilizar_xlsx(caminho)


def criar_csv(registros: list[dict[str, object]], caminho: Path) -> None:
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(["Data", "Conta Contábil", "Valor", "Tipo", "Histórico", "Filial"])
        for item in registros:
            escritor.writerow(
                [
                    item["data"].strftime("%d/%m/%Y"),
                    item["conta"],
                    valor_br(item["valor"]),
                    item["tipo"],
                    item["historico"],
                    item["filial"],
                ]
            )


def criar_xlsx_multiplas_abas(registros: list[dict[str, object]], caminho: Path) -> None:
    livro = Workbook()
    preparar_livro(livro)
    resumo = livro.active
    resumo.title = "Resumo"
    resumo.append(["Arquivo de aceitação anonimizado"])
    resumo.append(["Quantidade de movimentos", len(registros)])
    resumo.append(["Total de débitos", sum(item["valor"] for item in registros if item["tipo"] == "D")])
    resumo.append(["Total de créditos", sum(item["valor"] for item in registros if item["tipo"] == "C")])
    aba = livro.create_sheet("Movimentos")
    aba.append(["Dia", "Conta", "Montante", "Natureza", "Memo", "Unidade"])
    for item in registros:
        aba.append(
            [
                item["data"],
                item["conta"],
                float(item["valor"]),
                item["tipo"],
                item["historico"],
                item["filial"],
            ]
        )
    for celula in aba["A"][1:]:
        celula.number_format = "DD/MM/YYYY"
    for celula in aba["C"][1:]:
        celula.number_format = '#,##0.00'
    livro.save(caminho)
    estabilizar_xlsx(caminho)


def gerar() -> None:
    registros = ler_origem()
    DESTINO.mkdir(parents=True, exist_ok=True)
    arquivos = {
        "cap_anonimizada_delimitada.xlsx": criar_xlsx_delimitado,
        "cap_anonimizada_colunas.csv": criar_csv,
        "cap_anonimizada_multiplas_abas.xlsx": criar_xlsx_multiplas_abas,
    }
    manifesto: dict[str, object] = {
        "origem": "CAP real, anonimizada e não versionada",
        "linhas_por_arquivo": len(registros),
        "debitos": sum(1 for item in registros if item["tipo"] == "D"),
        "creditos": sum(1 for item in registros if item["tipo"] == "C"),
        "total_debitos": str(sum(item["valor"] for item in registros if item["tipo"] == "D")),
        "total_creditos": str(sum(item["valor"] for item in registros if item["tipo"] == "C")),
        "arquivos": {},
    }
    for nome, criador in arquivos.items():
        caminho = DESTINO / nome
        criador(registros, caminho)
        manifesto["arquivos"][nome] = {
            "bytes": caminho.stat().st_size,
            "sha256": hashlib.sha256(caminho.read_bytes()).hexdigest(),
        }
    (DESTINO / "manifesto.json").write_text(
        json.dumps(manifesto, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Fixtures criadas em {DESTINO}")


if __name__ == "__main__":
    gerar()
